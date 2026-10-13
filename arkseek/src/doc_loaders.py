# -*- coding: utf-8 -*-
"""
各种格式的文档读取器（知识库扩充：JC讲稿/PPT/文献）

知识库原来只读 .txt（两份课题报告是手动用pdfplumber转成txt放进去的）。扩充之后
要收录的格式：文献PDF（英文）、JC讲稿（Word，大部分中文，25.7.5那次是英文；
26.9.12那次是Markdown）、笔记脉络（Word/Markdown）、PPT（英文，内容主要在每页
正文里，演讲者备注很少）。每种格式一个读取函数，统一返回纯文本，交给
chunking.chunk_document() 按token数切片——业界做法（LangChain/LlamaIndex的
document loader）也是"每种格式一个loader，输出统一的文本+元数据"。

依赖：python-docx、python-pptx、pymupdf（用户本机已装）。
"""

import re
from typing import List

# 清洗规则（参考文献识别、页眉页脚、软连字符等）改动时加1：kb_sources.py按这个
# 版本号判断缓存的解析结果是否还能用
PARSER_VERSION = 1

# 论文PDF的清洗：按"文本块"（PyMuPDF的blocks：一个段落/标题/图注一块）处理。
#
# 一开始用 get_text(sort=True) 按版面行排序，真机检查发现双栏论文里会把"REFERENCES"
# 标题和另一栏同一行的正文粘在一起，按标题切参考文献完全失效（三篇测试论文一个字
# 都没切掉）。改成按块处理：块的默认顺序就是论文排版的阅读顺序（先左栏再右栏），
# 标题、段落各自是完整的一块。
#
# 参考文献列表：检索时基本只有坏处（满屏作者名/期刊名会跟问题里的专有名词撞车），
# 要去掉。但不能靠"从References标题切到结尾"——Cell系列参考文献后面还有STAR★Methods
# 正文，Nature系列的"References"标题常常跟第一条文献并在同一块里。所以逐块判断
# "这一块是不是参考文献条目"：以编号开头、含年份、再加上"et al."/页码范围/doi之一。
# 页眉页脚（"Nature Methods | Volume 22 | March 2025"这种）按"短、且在多页重复出现"去掉。
# 一条参考文献的开头："编号. 姓, 名缩写"（"33. Katsuda, T.,"、"\t2.\t Cusanovich, D. A."）。
# 要求编号后面紧跟"姓, 缩写"这个形状，方法部分"1. Cells were…"这类编号步骤不会误判。
# 检查块内任意一行（ArchR这类排版一个块会跨两条文献：上一条的尾巴+下一条的开头）。
_CITATION_START = re.compile(r"^\s*\d{1,3}\.\s+[A-Z][A-Za-z'’\-]+(?: [A-Z][A-Za-z'’\-]+)?,\s+[A-Z]\.")
_YEAR = re.compile(r"\b(19|20)\d{2}\b")
_CITATION_HINT = re.compile(r"et al\.|\d+[–-]\d+|doi|https?://|\bVol\.|\(\d{4}\)")
_REF_HEADING = re.compile(r"^\s*(REFERENCES|References)\s*$")


def _is_reference_block(txt: str) -> bool:
    if _REF_HEADING.match(txt):
        return True
    has_start = any(_CITATION_START.match(ln) for ln in txt.splitlines())
    return bool(has_start and _YEAR.search(txt) and _CITATION_HINT.search(txt))


def _clean_block(txt: str) -> str:
    txt = re.sub("­\\s*", "", txt)               # 软连字符（排版断词，"influ­ ences"）
    txt = re.sub(r"(\w)-\n(\w)", r"\1\2", txt)       # 行尾连字符断词接回去
    return re.sub(r"\s*\n\s*", " ", txt).strip()     # 块内硬换行换成空格


def read_pdf(path: str, drop_references: bool = True) -> str:
    import pymupdf
    from collections import Counter

    doc = pymupdf.open(path)
    pages = [[b[4] for b in page.get_text("blocks") if b[6] == 0] for page in doc]  # b[6]==0：文本块（不是图片）
    # 页眉页脚：短文本、在至少3页上一字不差地重复出现（数字换成#再比，页码不同也算同一个）
    norm = lambda t: re.sub(r"\d+", "#", t.strip())
    freq = Counter(norm(t) for blocks in pages for t in set(blocks) if len(t.strip()) < 100)
    kept = []
    for blocks in pages:
        for t in blocks:
            if not t.strip():
                continue
            if len(t.strip()) < 100 and freq[norm(t)] >= 3:
                continue
            if drop_references and _is_reference_block(t):
                continue
            kept.append(_clean_block(t))
    return "\n\n".join(kept)


def read_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(p.strip() for p in parts if p.strip())


def read_pptx(path: str) -> str:
    """每页一段："第N页：正文（演讲者备注：…）"。保留页码，回答时可以说"PPT第几页"。"""
    import pptx
    prs = pptx.Presentation(path)
    pages: List[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        body = " / ".join(t.replace("\n", " ") for t in texts)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if body or notes:
            pages.append(f"第{i}页：{body}" + (f"（演讲者备注：{notes}）" if notes else ""))
    return "\n\n".join(pages)


def read_markdown(path: str) -> str:
    with open(path, encoding="utf-8") as f:
        text = f.read()
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)  # 图片引用
    return text.strip()


def read_any(path: str) -> str:
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return read_pdf(path)
    if ext == "docx":
        return read_docx(path)
    if ext == "pptx":
        return read_pptx(path)
    if ext in ("md", "markdown"):
        return read_markdown(path)
    if ext == "txt":
        with open(path, encoding="utf-8") as f:
            return f.read()
    raise ValueError(f"不支持的格式：{path}")
